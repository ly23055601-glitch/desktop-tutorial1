#!/usr/bin/env python3
"""Parse one actual social-helper XLSX into provenance-preserving JSONL.

This parser never opens social URLs, updates a collection ledger, marks a work
successful, or reads media. Use the bundled Python runtime with openpyxl.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
import hashlib
import json
import math
import os
from pathlib import Path
import re
import sys
import tempfile
from urllib.parse import urlsplit, urlunsplit
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

PLATFORMS = {"douyin", "xiaohongshu", "bilibili"}
SPLITS = {"train", "holdout", "legacy"}
DOMAIN = {"douyin": "douyin.com", "xiaohongshu": "xiaohongshu.com", "bilibili": "bilibili.com"}
ID_PATTERN = {"douyin": re.compile(r"(?<!\d)\d{15,22}(?!\d)"),
              "xiaohongshu": re.compile(r"(?<![a-fA-F0-9])[a-fA-F0-9]{24}(?![a-fA-F0-9])"),
              "bilibili": re.compile(r"(?<![A-Za-z0-9])(?:BV[A-Za-z0-9]{10}|av\d+)(?![A-Za-z0-9])")}
COMMON_FIELDS = {
    "captured_at": ["采集时间"], "post_title": ["视频标题", "笔记标题"],
    "post_text": ["视频描述", "笔记内容"], "post_topics": ["视频话题", "笔记话题"],
    "post_tags": ["视频标签", "笔记标签"], "post_type": ["视频类型", "笔记类型"],
    "post_created_at": ["发布时间", "视频发布时间", "笔记发布时间"],
    "post_author_name": ["达人昵称", "博主昵称", "UP主名称"],
    "post_author_id": ["达人ID", "博主ID", "UP主ID"],
    "post_comment_count": ["视频评论量", "笔记评论量"],
    "post_likes": ["视频点赞量", "笔记点赞量"],
    "cover_url": ["视频封面链接", "笔记封面链接"],
    "media_url": ["视频文件链接", "笔记视频链接"], "audio_url": ["音频文件链接"],
}
COMMENT_FIELDS = {
    "comment_id": ["评论ID"], "text": ["评论内容"], "image_urls": ["评论图片链接"],
    "comment_likes": ["点赞量"], "created_at": ["评论时间"],
    "reported_reply_count": ["子评论数"], "author_name": ["用户名称"],
    "author_id": ["用户ID"], "root_comment_id": ["一级评论ID"],
    "root_comment_text": ["一级评论内容"], "root_author_name": ["一级评论用户名称"],
    "referenced_comment_id": ["引用的评论ID"], "referenced_comment_text": ["引用的评论内容"],
    "referenced_author_name": ["引用的用户名称"],
}
ID_FIELDS = {"work_id", "aid", "post_author_id", "comment_id", "author_id", "root_comment_id", "referenced_comment_id"}
MODEL_RE = re.compile(r"(?<![A-Za-z0-9])(?:(?:DJI|大疆)\s*)?(?:(?:Osmo\s*)?Pocket\s*(?P<full>[1-4])\s*(?P<suffix>Pro|P)?|(?P<short>p[1-4]p?|4p))(?![A-Za-z0-9])", re.IGNORECASE)
COMMON_LIMITATIONS = [
    "仅解析实际导出字段，不认证发布者或评论者的经历、持有型号和产品说法",
    "媒体链接只作来源记录；未读取图片、连续视频、音轨或完整对白",
    "导出行顺序不证明评论排名、置顶或直接回复关系",
    "解析完成不表示平台采集成功；本程序不修改ledger或作品完成状态",
]


class ExportError(ValueError):
    """Errors may include file coordinates and headers, never comment bodies."""


def json_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    return value


def work_ids(value, platform):
    if value is None or value == "":
        return set()
    text = str(value).strip()
    if "://" in text:
        parsed = urlsplit(text)
        if parsed.hostname != DOMAIN[platform] and not (parsed.hostname or "").endswith("." + DOMAIN[platform]):
            return set()
        text = parsed.path
    found = set(ID_PATTERN[platform].findall(text))
    return {x.lower() if platform == "xiaohongshu" else x for x in found}


def canonical_url(value):
    """Only local raw_fields retain source query tokens; public references do not."""
    if not isinstance(value, str):
        return None
    parts = urlsplit(value.strip())
    if parts.scheme not in {"https", "http"} or not parts.hostname:
        return None
    return urlunsplit((parts.scheme, parts.netloc, parts.path, "", ""))


def frozen_ledger(path, platform):
    items, lookup = [], defaultdict(list)
    for number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        if not line.strip():
            continue
        try:
            item = json.loads(line)
        except json.JSONDecodeError as exc:
            raise ExportError(f"ledger line {number}: invalid JSON") from exc
        if not isinstance(item, dict):
            raise ExportError(f"ledger line {number}: expected object")
        if item.get("platform") != platform:
            continue
        if item.get("split") not in SPLITS:
            raise ExportError(f"ledger line {number}: split must already be train/holdout/legacy")
        identifiers = set()
        for key in ("work_id", "workId", "canonical_url", "canonicalUrl", "url", "originalUrl"):
            identifiers.update(work_ids(item.get(key), platform))
        for alias in item.get("aliasRefs", []):
            identifiers.update(work_ids(alias, platform))
        if not identifiers:
            raise ExportError(f"ledger line {number}: no resolved platform work ID")
        label = item.get("id", item.get("inputOrder", item.get("input_order", number)))
        entry = {"input_id": str(label), "inputOrder": item.get("inputOrder", item.get("input_order")),
                 "sourceInputOrders": item.get("sourceInputOrders", []), "split": item["split"],
                 "ledger_line": number, "work_ids": sorted(identifiers)}
        items.append(entry)
        for identifier in identifiers:
            lookup[identifier].append(entry)
    if not items:
        raise ExportError("ledger has no entries for requested platform")
    if len({item["input_id"] for item in items}) != len(items):
        raise ExportError("ledger has duplicate input labels")
    for identifier, matches in lookup.items():
        if len(matches) > 1:
            # A duplicate is not silently assigned a split; upstream ledger owns dedup.
            raise ExportError(f"ledger work ID {identifier}: multiple inputs; reconcile ledger first")
    return items, lookup


def field_options(platform, kind, mapping=None):
    options = dict(COMMON_FIELDS)
    if kind == "comments":
        options.update(COMMENT_FIELDS)
    options.update({"work_id": ["笔记ID" if platform == "xiaohongshu" else "视频ID"],
                    "work_url": ["笔记链接" if platform == "xiaohongshu" else "视频链接"]})
    if platform == "bilibili":
        options["aid"] = ["视频AID"]
    if platform == "xiaohongshu" and kind == "posts":
        # Verified in the actual 2026-09-06 XHS batch-01 export, not inferred from UI labels.
        options["post_comment_count"] = ["笔记评论量", "评论量"]
        options["post_image_urls"] = ["笔记图片链接"]
        options["post_image_count"] = ["图片数量"]
        options["post_video_duration"] = ["笔记视频时长"]
    if mapping is not None:
        if not isinstance(mapping, dict) or not isinstance(mapping.get("fields"), dict):
            raise ExportError("mapping must contain an exact fields object")
        for key, header in mapping["fields"].items():
            if key not in options or not isinstance(header, str) or not header:
                raise ExportError("mapping contains unknown semantic field or invalid header")
            options[key] = [header]
    return options


def map_headers(headers, options, kind, where):
    if any(not isinstance(h, str) or not h.strip() for h in headers) or len(set(headers)) != len(headers):
        raise ExportError(f"{where}: empty, nontext or duplicate headers")
    result = {}
    for semantic, candidates in options.items():
        present = [candidate for candidate in candidates if candidate in headers]
        if len(present) > 1:
            raise ExportError(f"{where}: multiple candidate columns for {semantic}; supply exact --mapping")
        if present:
            result[semantic] = present[0]
    if not ({"work_id", "work_url"} & result.keys()):
        raise ExportError(f"{where}: missing work identity column")
    if kind == "comments" and not {"comment_id", "text"}.issubset(result):
        raise ExportError(f"{where}: missing comment ID or text header")
    if kind == "posts" and "评论ID" in headers:
        raise ExportError(f"{where}: comment table cannot be parsed as posts")
    if kind == "posts" and not ({"post_text", "post_title"} & result.keys()):
        raise ExportError(f"{where}: missing post title/body header")
    if len(set(result.values())) != len(result):
        raise ExportError(f"{where}: one column mapped to multiple meanings")
    return result


def id_string(cell):
    value = cell.value
    if value is None or value == "":
        return None, None
    if isinstance(value, str):
        return value.strip(), None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return str(json_value(value)), "non_text_identifier_type"
    if isinstance(value, float) and (not math.isfinite(value) or not value.is_integer()):
        return str(value), "fractional_or_nonfinite_identifier"
    text = str(int(value))
    return text, "numeric_identifier_over_15_digits" if len(text.lstrip("-")) > 15 else None


def model_evidence(fields):
    evidence = []
    pocket_context = any(isinstance(value, str) and re.search(r"Pocket", value, re.IGNORECASE) for value in fields.values())
    for field, value in fields.items():
        if not isinstance(value, str):
            continue
        for match in MODEL_RE.finditer(value):
            if match["short"]:
                short = match["short"].lower()
                suffix = short.endswith("p")
                digit = short[1] if short.startswith("p") else short[0]
            else:
                digit, suffix = match["full"], bool(match["suffix"])
            if suffix and digit != "4":
                continue
            evidence.append({"model_id": f"pocket_{digit}" + ("p" if suffix else ""),
                             "field": field, "exact_model_token": match.group(0),
                             "basis": "ambiguous_shorthand_without_pocket_context" if match["short"] and not pocket_context else "comment_text" if field == "text" else "exported_post_text_or_tag"})
    candidates = sorted({e["model_id"] for e in evidence})
    explicit = len(candidates) == 1 and any(e["basis"] != "ambiguous_shorthand_without_pocket_context" for e in evidence)
    return {"model_id": candidates[0] if explicit else "unknown", "model_candidates": candidates,
            "model_basis": "single_explicit_exported_model" if explicit else "mixed_explicit_models" if len(candidates) > 1 else "ambiguous_shorthand" if candidates else "not_established_in_export",
            "model_evidence": evidence}


def images(value):
    if isinstance(value, (list, tuple)):
        return [str(item) for item in value if item]
    if isinstance(value, str):
        try:
            parsed = json.loads(value)
            if isinstance(parsed, list):
                return [str(item) for item in parsed if item]
        except json.JSONDecodeError:
            pass
        return re.findall(r"https?://[^\s<>\"'，,]+", value)
    return []


def time_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    # Preserve wall time and absence of timezone. Never substitute ingestion time.
    return value if isinstance(value, str) and value.strip() else None


def parse_export(platform, kind, source, ledger_path, mapping=None):
    source = source.resolve()
    items, lookup = frozen_ledger(ledger_path, platform)
    options = field_options(platform, kind, mapping)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    workbook = load_workbook(source, read_only=True, data_only=False)
    output, schemas, blank_rows = [], [], 0
    all_seen = {}
    try:
        for sheet in workbook:
            iterator = sheet.iter_rows()
            header_cells = next(iterator, ())
            headers = [cell.value for cell in header_cells]
            if not headers or all(h is None for h in headers):
                continue
            # Discard unused trailing columns only; interior missing headers are errors.
            while headers and headers[-1] is None:
                headers.pop()
            fields = map_headers(headers, options, kind, f"sheet {sheet.title}")
            columns = {semantic: {"header": header, "column": get_column_letter(headers.index(header) + 1)} for semantic, header in fields.items()}
            schemas.append({"sheet": sheet.title, "header_row": 1, "headers": headers, "semantic_columns": columns,
                            "unmapped_headers": [h for h in headers if h not in fields.values()]})
            for row_number, cells in enumerate(iterator, 2):
                if all(cell.value is None for cell in cells):
                    blank_rows += 1
                    continue
                if any(cell.value is not None for cell in cells[len(headers):]):
                    raise ExportError(f"sheet {sheet.title} row {row_number}: data under a missing header")
                row_cells = dict(zip(headers, cells))
                raw = {header: json_value(cell.value) for header, cell in row_cells.items()}
                normalized = {semantic: row_cells[header].value for semantic, header in fields.items()}
                issues = []
                formula_columns = [get_column_letter(i + 1) for i, cell in enumerate(cells) if cell.data_type == "f"]
                if formula_columns:
                    issues.append({"code": "formula_is_not_exported_value", "columns": formula_columns})
                unsafe_ids = []
                for semantic in fields.keys() & ID_FIELDS:
                    normalized[semantic], problem = id_string(row_cells[fields[semantic]])
                    if problem:
                        unsafe_ids.append(semantic)
                        issues.append({"code": problem, "field": semantic})
                observed_ids, matchable_ids = set(), set()
                for semantic in ("work_id", "work_url"):
                    value = normalized.get(semantic)
                    identifiers = work_ids(value, platform)
                    observed_ids.update(identifiers)
                    if semantic not in unsafe_ids and (semantic not in fields or row_cells[fields[semantic]].data_type != "f"):
                        matchable_ids.update(identifiers)
                    if value not in (None, "") and not identifiers:
                        issues.append({"code": "unrecognized_work_identity", "field": semantic})
                if platform == "bilibili" and normalized.get("aid") and str(normalized["aid"]).isdigit():
                    observed_ids.add("av" + normalized["aid"])
                    if "aid" not in unsafe_ids:
                        matchable_ids.add("av" + normalized["aid"])
                matched_entries = {entry["input_id"]: entry for identifier in matchable_ids for entry in lookup.get(identifier, [])}
                if len(matched_entries) == 1:
                    match = next(iter(matched_entries.values()))
                    match_status = "matched"
                else:
                    match = None
                    match_status = "ambiguous" if len(matched_entries) > 1 else "unsafe_identity" if not matchable_ids and unsafe_ids else "unmatched"
                    issues.append({"code": "ledger_" + match_status})
                if not observed_ids:
                    issues.append({"code": "work_identity_missing"})
                # IDs from URL and ID columns must agree, except the same row's BVID/AID pair.
                primary = work_ids(normalized.get("work_id"), platform)
                url_ids = work_ids(normalized.get("work_url"), platform)
                verified_aid_pair = platform == "bilibili" and normalized.get("aid") and ("av" + normalized["aid"]) in url_ids and any(x.startswith("BV") for x in primary)
                if primary and url_ids and primary.isdisjoint(url_ids) and "work_id" not in unsafe_ids and not verified_aid_pair:
                    issues.append({"code": "work_id_url_disagree"})
                    match_status, match = "ambiguous", None
                chosen_work = next(iter(primary), None) if "work_id" not in unsafe_ids else None
                chosen_work = chosen_work or next(iter(url_ids), None) or next(iter(primary), None)
                split = match["split"] if match else "holdout" if any(e["split"] == "holdout" for e in matched_entries.values()) else None
                post_fields = {key: json_value(value) for key, value in normalized.items() if key.startswith("post_")}
                if kind == "comments":
                    model_fields = {k: normalized.get(k) for k in ("post_title", "post_text", "post_topics", "post_tags", "text")}
                else:
                    model_fields = {k: normalized.get(k) for k in ("post_title", "post_text", "post_topics", "post_tags")}
                model = model_evidence(model_fields)
                common = {"platform": platform, "kind": kind, "material_kind": "comment" if kind == "comments" else "public_post", "work_id": chosen_work,
                          "canonical_url": canonical_url(normalized.get("work_url")), "work_id_candidates": sorted(observed_ids),
                          "input_id": match["input_id"] if match else None,
                          "inputOrder": match["inputOrder"] if match else None,
                          "sourceInputOrders": match["sourceInputOrders"] if match else [],
                          "split": split, "ledger_match": match_status,
                          "split_basis": "frozen_ledger_match" if match else "conservative_holdout_isolation_for_ambiguous_identity" if split == "holdout" else "unassigned",
                          "source_path": str(source), "source_sha256": digest,
                          "source_locator": {"sheet": sheet.title, "row": row_number},
                          "source_columns": columns, "source_cell_types": {header: {"type": cell.data_type, "number_format": cell.number_format} for header, cell in row_cells.items()},
                          "raw_fields": raw, "post_fields": post_fields,
                          "collected_at": time_value(normalized.get("captured_at")),
                          "timestamp_timezone": "not_recorded_unless_explicit_in_export_value", **model,
                          "limitations": list(COMMON_LIMITATIONS)}
                if not common["collected_at"]:
                    issues.append({"code": "collection_time_missing"})
                if kind == "comments":
                    comment_id = normalized.get("comment_id")
                    text = normalized.get("text")
                    if text is not None and not isinstance(text, str):
                        issues.append({"code": "comment_text_nontext_cell"})
                        text = str(json_value(text))
                    namespace = platform + ":"
                    root_id, parent_id = normalized.get("root_comment_id"), normalized.get("referenced_comment_id")
                    if "root_comment_id" not in fields and "referenced_comment_id" not in fields:
                        role = "unknown"
                    else:
                        role = "root" if root_id in (None, comment_id) and parent_id is None else "reply"
                    empty_text = text is None or not text.strip()
                    if role == "reply" and not parent_id:
                        common["limitations"].append("缺少直接父句ID，一级评论ID只表示根线程，不据相邻行或@文字补父句")
                    if role == "unknown":
                        common["limitations"].append("导出缺少线程关系列，无法确认本行为根评论或回复")
                    if empty_text:
                        common["limitations"].append("评论文字为空；图片链接即使存在也未读取，不用本行提炼文字内容")
                    if not comment_id:
                        issues.append({"code": "comment_id_missing"})
                    common.update({"id": namespace + comment_id if comment_id else None, "comment_id": comment_id,
                                   "text": text, "text_status": "absent_in_export" if empty_text else "available",
                                   "image_urls": images(normalized.get("image_urls")),
                                   "parent_id": namespace + parent_id if parent_id else None,
                                   "root_id": namespace + root_id if root_id and root_id != comment_id else None,
                                   "thread_role": role, "parent_basis": "exported_reference_id" if parent_id else "not_provided",
                                   "referenced_comment_text": normalized.get("referenced_comment_text"),
                                   "root_comment_text": normalized.get("root_comment_text"),
                                   "author_id": normalized.get("author_id"), "author_name": normalized.get("author_name"),
                                   "created_at": time_value(normalized.get("created_at")),
                                   "comment_likes": json_value(normalized.get("comment_likes")),
                                   "reported_reply_count": json_value(normalized.get("reported_reply_count")),
                                   "like_scope": "exported_comment_likes_only",
                                   "rank_status": "export_order_unverified", "pinned_status": "not_established",
                                   "read_scope": "exported_comment_and_embedded_post_text_only; no_media_review"})
                else:
                    title, text = normalized.get("post_title"), normalized.get("post_text")
                    if not title and not text:
                        issues.append({"code": "post_title_and_body_empty"})
                    common.update({"id": f"{platform}:{chosen_work}" if chosen_work else None,
                                   "title": title, "text": text, "author_id": normalized.get("post_author_id"),
                                   "author_name": normalized.get("post_author_name"),
                                   "author_identity_status": "exported_id" if normalized.get("post_author_id") else "display_name_only" if normalized.get("post_author_name") else "not_in_export",
                                   "created_at": time_value(normalized.get("post_created_at")),
                                   "cover_url": normalized.get("cover_url"), "media_url": normalized.get("media_url"), "audio_url": normalized.get("audio_url"),
                                   "read_scope": "exported_post_text_metadata_only; media_urls_not_reviewed"})
                identifier = common["id"]
                if identifier in all_seen:
                    issues.append({"code": "duplicate_id_retained", "previous": all_seen[identifier]["source_locator"]})
                    if all_seen[identifier]["work_id"] != chosen_work:
                        issues.append({"code": "comment_id_under_different_work"})
                elif identifier:
                    all_seen[identifier] = common
                common["normalization_issues"] = issues
                common["normalization_status"] = "needs_review" if issues else "parsed"
                output.append(common)
    finally:
        workbook.close()
    # Resolve only explicit parent IDs across the complete file, never by row order.
    by_id = {r["id"]: r for r in output if r.get("id")}
    if kind == "comments":
        for record in output:
            parent = by_id.get(record.get("parent_id"))
            if parent and parent.get("work_id") == record.get("work_id"):
                record["parent_source_locator"] = parent["source_locator"]
                record["parent_lookup"] = "present_same_work"
            else:
                record["parent_source_locator"] = None
                record["parent_lookup"] = "different_work" if parent else "not_in_file" if record.get("parent_id") else "not_provided"
                if parent:
                    record["normalization_issues"].append({"code": "parent_id_under_different_work"})
                    record["normalization_status"] = "needs_review"
            if record.get("parent_id") and record.get("referenced_comment_text"):
                record["parent_quote_source_locator"] = {**record["source_locator"], "column": record["source_columns"]["referenced_comment_text"]["column"]}
    summaries = []
    for entry in items:
        rows = [r for r in output if r.get("input_id") == entry["input_id"]]
        summaries.append({**entry, "matched_rows": len(rows), "unique_row_ids": len({r["id"] for r in rows if r.get("id")}),
                          "needs_review_rows": sum(r["normalization_status"] == "needs_review" for r in rows),
                          "root_rows": sum(r.get("thread_role") == "root" for r in rows),
                          "reply_rows": sum(r.get("thread_role") == "reply" for r in rows),
                          "explicit_parents_in_file": sum(r.get("parent_lookup") == "present_same_work" for r in rows)})
    report = {"platform": platform, "kind": kind, "normalization_status": "needs_review" if any(r["normalization_status"] == "needs_review" for r in output) else "parsed" if output else "empty",
              "collection_status": "not_decided_by_parser", "ledger_updated": False,
              "source_path": str(source), "source_sha256": digest, "ledger_path": str(ledger_path.resolve()),
              "parsed_at": datetime.now(timezone.utc).isoformat(), "schemas": schemas,
              "data_rows": len(output), "blank_rows": blank_rows,
              "matched_rows": sum(r["ledger_match"] == "matched" for r in output),
              "unmatched_or_ambiguous_rows": sum(r["ledger_match"] != "matched" for r in output),
              "by_split": dict(Counter(r.get("split") or "unassigned" for r in output)),
              "issue_counts": dict(Counter(i["code"] for r in output for i in r["normalization_issues"])),
              "inputs": summaries, "inputs_without_export_rows": [e["input_id"] for e in summaries if not e["matched_rows"]],
              "holdout_text_logged": False, "limitations": COMMON_LIMITATIONS}
    return output, report


def atomic_write(path, text):
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix=path.name + ".", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as handle:
            handle.write(text)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--platform", required=True, choices=sorted(PLATFORMS))
    parser.add_argument("--kind", required=True, choices=["posts", "comments"])
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--ledger", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path, help="Explicit normalized JSONL path; sibling .reconcile.json is also written")
    parser.add_argument("--mapping", type=Path, help='Optional verified header override: {"fields":{"semantic":"exact header"}}')
    parser.add_argument("--overwrite", action="store_true", help="Replace explicitly selected output files after reprocessing")
    args = parser.parse_args(argv)
    report_path = args.output.with_suffix(".reconcile.json")
    try:
        if args.source.suffix.lower() != ".xlsx":
            raise ExportError("source must be an actual .xlsx export")
        if any(p.resolve() in {args.source.resolve(), args.ledger.resolve()} for p in (args.output, report_path)):
            raise ExportError("output must not overwrite source or ledger")
        if not args.overwrite and (args.output.exists() or report_path.exists()):
            raise ExportError("output already exists; choose another path or explicitly use --overwrite")
        ledger_digest = hashlib.sha256(args.ledger.read_bytes()).hexdigest()
        ledger_identity, _ = frozen_ledger(args.ledger, args.platform)
        mapping = json.loads(args.mapping.read_text(encoding="utf-8")) if args.mapping else None
        records, report = parse_export(args.platform, args.kind, args.source, args.ledger, mapping)
        latest_identity, _ = frozen_ledger(args.ledger, args.platform)
        if latest_identity != ledger_identity:
            raise ExportError("target work identities or frozen splits changed during parsing; rerun")
        report["ledger_sha256"] = ledger_digest
        report["ledger_metadata_changed_during_parse"] = hashlib.sha256(args.ledger.read_bytes()).hexdigest() != ledger_digest
        atomic_write(args.output, "".join(json.dumps(r, ensure_ascii=False) + "\n" for r in records))
        atomic_write(report_path, json.dumps(report, ensure_ascii=False, indent=2) + "\n")
        summary = {key: report[key] for key in ("platform", "kind", "normalization_status", "collection_status", "data_rows", "matched_rows", "unmatched_or_ambiguous_rows", "by_split", "issue_counts", "inputs_without_export_rows", "ledger_updated", "holdout_text_logged")}
        summary.update(output=str(args.output.resolve()), reconciliation=str(report_path.resolve()))
        print(json.dumps(summary, ensure_ascii=False))
        return 0 if report["normalization_status"] == "parsed" else 1
    except (ExportError, OSError, json.JSONDecodeError, BadZipFile, InvalidFileException) as exc:
        print(json.dumps({"normalization_status": "error", "error": str(exc), "holdout_text_logged": False, "ledger_updated": False}, ensure_ascii=False), file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
